# -*- coding: utf-8 -*-
import openpyxl
from pathlib import Path

EXCEL_PATH = Path(r"c:\Users\Вячеслав\Desktop\Контур\Май\301 гвозди.xlsx")
HTACCESS_PATH = Path(r"E:\Работа\OSPanel\domains\privarka\.htaccess")
REDIRECTS_PATH = Path(r"E:\Работа\OSPanel\domains\privarka\local\redirects\gvozdi_301.htaccess")

OLD_PREFIX = "https://www.privarka-k97.ru/krepezh/privarnoy_krepyezh/gvozdi-dlya-krepleniya-izolyatsii/"
NEW_PREFIX = "https://privarka-k97.ru/krepezh/gvozdi-dlya-krepleniya-izolyatsii/"
OLD_PATH_PREFIX = "/krepezh/privarnoy_krepyezh/gvozdi-dlya-krepleniya-izolyatsii/"
NEW_PATH_PREFIX = "/krepezh/gvozdi-dlya-krepleniya-izolyatsii/"


def url_to_path(url: str) -> str:
    path = url.split("://", 1)[1]
    path = path.split("/", 1)[1]
    return "/" + path if not path.startswith("/") else "/" + path.lstrip("/")


def path_to_old_url(path: str) -> str:
    if path.startswith(NEW_PATH_PREFIX):
        suffix = path[len(NEW_PATH_PREFIX):]
        return OLD_PREFIX + suffix
    if path.startswith(OLD_PATH_PREFIX):
        suffix = path[len(OLD_PATH_PREFIX):]
        return NEW_PREFIX + suffix
    raise ValueError(f"Unexpected path: {path}")


def path_to_new_url(path: str) -> str:
    if path.startswith(OLD_PATH_PREFIX):
        suffix = path[len(OLD_PATH_PREFIX):]
        return NEW_PREFIX + suffix
    if path.startswith(NEW_PATH_PREFIX):
        suffix = path[len(NEW_PATH_PREFIX):]
        return NEW_PREFIX + suffix
    raise ValueError(f"Unexpected path: {path}")


def main() -> None:
    wb = openpyxl.load_workbook(EXCEL_PATH)
    ws = wb.active

    ws.cell(row=1, column=1, value="Старый URL")
    ws.cell(row=1, column=2, value="Новый URL")

    pairs: list[tuple[str, str]] = []

    for row_idx in range(2, ws.max_row + 1):
        value = ws.cell(row=row_idx, column=1).value
        if not value or not isinstance(value, str):
            continue

        value = value.strip()
        if value.startswith(OLD_PREFIX):
            old_url = value
            new_url = NEW_PREFIX + value[len(OLD_PREFIX):]
        elif value.startswith(NEW_PREFIX):
            new_url = value
            old_url = OLD_PREFIX + value[len(NEW_PREFIX):]
        elif value.startswith("http"):
            old_path = url_to_path(value)
            if OLD_PATH_PREFIX in old_path:
                new_path = old_path.replace(OLD_PATH_PREFIX, NEW_PATH_PREFIX, 1)
                old_url = "https://www.privarka-k97.ru" + old_path
                new_url = "https://privarka-k97.ru" + new_path
            elif NEW_PATH_PREFIX in old_path:
                new_path = old_path
                old_path = old_path.replace(NEW_PATH_PREFIX, OLD_PATH_PREFIX, 1)
                old_url = "https://www.privarka-k97.ru" + old_path
                new_url = "https://privarka-k97.ru" + new_path
            else:
                raise ValueError(f"Row {row_idx}: unknown URL format: {value}")
        else:
            raise ValueError(f"Row {row_idx}: unknown value: {value}")

        ws.cell(row=row_idx, column=1, value=old_url)
        ws.cell(row=row_idx, column=2, value=new_url)

        old_path = url_to_path(old_url)
        new_path = url_to_path(new_url)
        pairs.append((old_path, new_path))

    wb.save(EXCEL_PATH)

    REDIRECTS_PATH.parent.mkdir(parents=True, exist_ok=True)
    lines = [
        "# 301 редиректы: гвозди для крепления изоляции (перенос раздела)",
        "# Сгенерировано из Excel, относительные пути без домена",
        "<IfModule mod_alias.c>",
    ]
    for old_path, new_path in pairs:
        lines.append(f"Redirect 301 {old_path} {new_path}")
    lines.append("</IfModule>")
    REDIRECTS_PATH.write_text("\n".join(lines) + "\n", encoding="utf-8")

    htaccess = HTACCESS_PATH.read_text(encoding="utf-8")
    marker_start = "\t# 301: раздел «Гвозди для крепления изоляции»"
    marker_end = "\tRewriteCond %{REQUEST_FILENAME} !-f"

    if marker_start in htaccess:
        before = htaccess.split(marker_start)[0]
        after = htaccess.split(marker_end, 1)[1]
        new_block = (
            "\t# 301: раздел «Гвозди для крепления изоляции» — общее правило (эквивалентно "
            f"{len(pairs)} отдельным редиректам из Excel)\n"
            "\tRewriteRule ^krepezh/privarnoy_krepyezh/gvozdi-dlya-krepleniya-izolyatsii/?(.*)$ "
            "/krepezh/gvozdi-dlya-krepleniya-izolyatsii/$1 [R=301,L]\n\n"
            "\tRewriteCond %{REQUEST_FILENAME} !-f"
        )
        htaccess = before + new_block + after
        HTACCESS_PATH.write_text(htaccess, encoding="utf-8")

    print(f"Excel rows updated: {len(pairs)}")
    print(f"Redirects file: {REDIRECTS_PATH}")
    print(f"Sample old: {pairs[0][0]}")
    print(f"Sample new: {pairs[0][1]}")


if __name__ == "__main__":
    main()
